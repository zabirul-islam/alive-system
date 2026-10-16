"""Repair an expert rater's returned CSV whose long text columns were split by unquoted commas or newlines.
Ratings are recovered as 10-field blocks matching the rating vocabulary; each block is assigned to the item whose
Question text (from our master workbook) precedes it in the record, never to a bare number. Ratings are written
into a fresh copy of the master template (dropdowns kept).
Usage: python repair_rater_return.py returned.csv N     -> revision/expert/ALIVE_expert_review_raterN.xlsx
"""
import csv, json, re, sys
from pathlib import Path
from openpyxl import load_workbook

R = Path(__file__).resolve().parents[1]; E = R / "expert"
src, n = Path(sys.argv[1]), int(sys.argv[2])
master = E / "_master_template.xlsx" if (E / "_master_template.xlsx").exists() else E / "ALIVE_expert_review_rater1.xlsx"
wb = load_workbook(master); ws = wb["Items"]
hdr = [c.value for c in ws[1]]; RC = hdr.index("Q answerable from lecture? (Y/N)")  # 0-based first rating column
norm = lambda s: re.sub(r"\s+", " ", str(s or "")).strip()
Q = {int(ws.cell(r, 1).value): norm(ws.cell(r, 5).value) for r in range(2, ws.max_row + 1)}
ERR = {"none", "conceptual", "mathematical", "causal", "numerical", "other"}
PAT = [{"Y", "N"}, {"Y", "P", "N"}, {"Y", "N"}, {"C", "P", "I"}, ERR, {"1", "2", "3", "4", "5"}, {"C", "P", "I"}, ERR,
       {"1", "2", "3", "4", "5"}, {"A", "B", "Tie"}]
fit = lambda blk: len(blk) == 10 and all(v in s for v, s in zip(blk, PAT))

recs = list(csv.reader(open(src, encoding="utf-8-sig", newline="")))[1:]
got, comments, flags = {}, {}, []
for rec in recs:
    f = [x.strip() for x in rec]
    # positions where a known Question text occurs in this record (own item and any merged item); questions with
    # commas may themselves be split, so match on the comma-joined, whitespace-free record text
    joined, owner = "", []
    for j, x in enumerate(f):
        t = re.sub(r"\s", "", x) + ","; joined += t; owner += [j] * len(t)
    anchors = []
    for it, q in Q.items():
        qs = re.sub(r"\s", "", q)
        k = joined.find(qs + ",")
        if qs and k >= 0:
            anchors.append((owner[k + len(qs)], it))  # field where the question ends
    anchors.sort()
    blocks = [j for j in range(len(f) - 9) if fit(f[j:j + 10])]
    for k, (pos, it) in enumerate(anchors):
        nxt = anchors[k + 1][0] if k + 1 < len(anchors) else len(f)
        bs = [b for b in blocks if pos < b < nxt]
        if len(bs) != 1:
            flags.append((it, f"{len(bs)} rating blocks after its question")); continue
        if it in got:
            flags.append((it, "duplicate")); continue
        got[it] = f[bs[0]:bs[0] + 10]
        tail = []
        for x in f[bs[0] + 10:nxt]:  # comment runs until the next merged item's Item id (a bare number) or record end
            if x.isdigit():
                break
            tail.append(x)
        # an unmatched quote in a comment can swallow the following CSV line; cut at the next "<newline><Item>," row start
        comments[it] = re.split(r"\n\s*\d{1,3},", ", ".join(t for t in tail if t))[0].strip()
    if not anchors:
        flags.append((rec[0] if rec else "?", "no known question text in record"))
missing = sorted(set(Q) - set(got))
for r in range(2, ws.max_row + 1):
    it = int(ws.cell(r, 1).value)
    if it in got:
        for j, v in enumerate(got[it]):
            ws.cell(r, RC + 1 + j).value = int(v) if v.isdigit() else v
        ws.cell(r, RC + 11).value = comments.get(it, "")
out = E / f"ALIVE_expert_review_rater{n}.xlsx"; wb.save(out)
rep = {"source": str(src), "recovered": len(got), "missing": missing, "flags": flags}
(E / "returned").mkdir(exist_ok=True); json.dump(rep, open(E / f"returned/repair_report_rater{n}.json", "w"), indent=1); print(json.dumps(rep, indent=1))

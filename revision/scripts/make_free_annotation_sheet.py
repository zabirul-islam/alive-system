"""Second relevance annotation for the 50 retrieval questions, same procedure as the first author's original labels:
for each question the annotator picks the relevant slides freely from the slide deck (All_Lectures_Slides.pdf).
No candidate list, no narration text, no lecture hint, no pause time. Questions in the order of retrieval_gt.json.
Output: revision/annotation/retrieval_relevance_annotator2.xlsx
"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

R = Path(__file__).resolve().parents[1]
GT = json.load(open(R.parent / "alive-eval-bundle/raw/retrieval_gt.json"))
wb = Workbook(); ws = wb.active; ws.title = "Questions"
hdr = ["Q", "Question", "Relevant slides (e.g. 13-27; 13-29)", "Minutes spent", "Notes (optional)"]
ws.append(hdr)
for c in ws[1]:
    c.font = Font(bold=True)
for i, g in enumerate(GT, 1):
    ws.append([i, g["question"], "", "", ""])
for col, w in zip("ABCDE", [5, 70, 40, 14, 40]):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"
out = R / "annotation/retrieval_relevance_annotator2.xlsx"; wb.save(out); print("wrote", out, len(GT), "questions")

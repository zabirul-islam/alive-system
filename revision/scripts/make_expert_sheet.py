"""R1-3: build the blinded expert-review workbook (100 test items, 2 raters) + private answer key.

Stratified by difficulty (proportional), seed 42. Each item shows the lecture-restricted context,
question, reference answer, and the Base and ALIVE answers (lecture-restricted condition) as
Answer A / Answer B in randomised order. The key mapping A/B -> model is written to a separate file
that raters must not see.
Output: revision/expert/ALIVE_expert_review_rater{1,2}.xlsx, revision/expert/KEY_do_not_share.json
"""
import json
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

R = Path(__file__).resolve().parents[1]
C = [json.loads(l) for l in open(R / "results/contexts_2x2.jsonl")]
G = {m: {r["i"]: r for r in map(json.loads, open(R / f"results/gen/gen_{m}.jsonl"))} for m in ["base", "alive"]}
rng = np.random.default_rng(42)
pick = []
for d, k in [("basic", 20), ("intermediate", 40), ("advanced", 40)]:
    ids = [c["i"] for c in C if c["difficulty"] == d]
    pick += sorted(rng.choice(ids, k, replace=False).tolist())
rows, key = [], {}
for n, i in enumerate(pick, 1):
    c = C[i]; flip = bool(rng.integers(2))
    a, b = ("alive", "base") if flip else ("base", "alive")
    key[n] = {"i": i, "A": a, "B": b}
    rows.append({"Item": n, "Difficulty": c["difficulty"], "Lecture": c["lecture"], "Lecture context": c["ctx_lecture"],
                 "Question": c["question"], "Reference answer": c["reference"],
                 "Answer A": G[a][i]["pred_lecture"], "Answer B": G[b][i]["pred_lecture"],
                 "Q answerable from lecture? (Y/N)": "", "Reference correct? (Y/N)": "", "Context supports reference? (Y/N)": "",
                 "A correctness (C/P/I)": "", "A error type": "", "A clarity (1-5)": "",
                 "B correctness (C/P/I)": "", "B error type": "", "B clarity (1-5)": "",
                 "Preference (A/B/Tie)": "", "Comment": ""})
guide = pd.DataFrame({"Field": [
    "Before you start", "Blinding", "Sources",
    "Q answerable from lecture?", "Reference correct?", "Context supports reference?", "Correctness",
    "Error type", "Clarity", "Preference", "Comment", "Return"], "Instruction": [
    "Read RATER_GUIDE.pdf. Work in the Items sheet and edit only the rating columns (Q answerable ... Comment).",
    "Answers A and B come from two AI systems in random order per item. Do not try to identify them; do not discuss items with the other rater.",
    "Use your expertise, the lecture context, and the course slides (lecture_slides/*.pdf). Do not use AI chatbots or web search.",
    "Y if the question can be answered from this course's lecture material; N otherwise.",
    "Y if the reference answer is factually and conceptually correct (it was written by an AI model and may be wrong).",
    "Y if the lecture context shown in the row contains the information needed for the reference answer.",
    "C = correct; P = correct but incomplete or with a minor inaccuracy; I = substantive error or does not answer. Judge against the lecture and your expertise, NOT against the reference.",
    "Most serious error: none / conceptual / mathematical / causal / numerical / other (off-topic, empty, truncated, self-contradictory).",
    "For an undergraduate: 5 clear and well organised; 4 minor awkwardness; 3 understandable but vague or wordy; 2 hard to follow; 1 confusing or misleading.",
    "Which answer would you rather a student received? Correctness first, then clarity. A, B, or Tie.",
    "Optional; note anything unclear or borderline.",
    "Save as ALIVE_expert_review_raterN_<initials>.xlsx and email it back."]})
out = R / "expert"; out.mkdir(exist_ok=True)
for r in [1, 2]:
    with pd.ExcelWriter(out / f"ALIVE_expert_review_rater{r}.xlsx") as w:
        guide.to_excel(w, sheet_name="Instructions", index=False)
        pd.DataFrame(rows).map(lambda v: ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v).to_excel(w, sheet_name="Items", index=False)
# usability: dropdowns, wrapped text, frozen header, column widths
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation
for r in [1, 2]:
    f = out / f"ALIVE_expert_review_rater{r}.xlsx"; wb = load_workbook(f); ws = wb["Items"]
    widths = {"A": 6, "B": 12, "C": 11, "D": 60, "E": 35, "F": 40, "G": 40, "H": 40}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for col in "IJKLMNOPQR":
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["S"].width = 30
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "E2"
    n = ws.max_row
    for cols, opts in [("IJK", '"Y,N"'), ("LO", '"C,P,I"'), ("MP", '"none,conceptual,mathematical,causal,numerical,other"'),
                       ("NQ", '"1,2,3,4,5"'), ("R", '"A,B,Tie"')]:
        dv = DataValidation(type="list", formula1=opts, allow_blank=True); ws.add_data_validation(dv)
        for col in cols:
            dv.add(f"{col}2:{col}{n}")
    wi = wb["Instructions"]; wi.column_dimensions["A"].width = 28; wi.column_dimensions["B"].width = 110
    for row in wi.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(f)
json.dump(key, open(out / "KEY_do_not_share.json", "w"), indent=1)
print("items", len(rows), "A=alive share", np.mean([v["A"] == "alive" for v in key.values()]))
